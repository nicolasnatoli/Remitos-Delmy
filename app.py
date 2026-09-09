import hashlib
import io
import mimetypes
import os
import re
import secrets
import tempfile
import calendar
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from dotenv import load_dotenv
from flask import (Flask, Response, abort, jsonify, redirect,
                    render_template, request, session, url_for)
from werkzeug.utils import secure_filename

load_dotenv()

import auth
import database as db
import email_import
import ocr_parser
from etiquetas_pdf import generar_pdf_etiquetas, generar_pdf_etiquetas_ubicaciones

ALLOWED_EXT = {".pdf", ".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}


def _get_or_create_secret():
    # Vercel no tiene disco persistente para guardar una clave generada acá
    # (y el proyecto queda como solo lectura en runtime) — hace falta
    # configurar SECRET_KEY como variable de entorno. En local, si no está
    # configurada, se genera una en memoria para esa sola ejecución (avisando
    # que las sesiones no van a sobrevivir un reinicio).
    env_key = os.environ.get("SECRET_KEY")
    if env_key:
        return env_key.encode()
    print("AVISO: falta SECRET_KEY en el entorno — se genera una temporal, "
          "las sesiones no van a persistir entre reinicios del servidor.")
    return secrets.token_bytes(32)


app = Flask(__name__)
app.secret_key = _get_or_create_secret()
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB
app.config["TEMPLATES_AUTO_RELOAD"] = True  # los .html se recargan solos, sin reiniciar el server

db.init_db()
auth.ensure_default_admin()


def _estado_detallado(r: dict) -> str:
    if r["estado"] == "pendiente":
        return "Pend. recepción"
    if r.get("articulos_pendientes"):
        return "Falta cargar artículos"
    if r["ingresado_sistema_en"]:
        return "Completado"
    if r["tipo_control"] == "bulto":
        return "Recepcionado x bulto — pend. ingreso OC"
    if r["tipo_control"] == "articulo":
        return "Recepcionado x artículo — pend. ingreso OC"
    return "Pend. ingreso OC"


def _tag_clase(r: dict) -> str:
    # La recepción (confirmar comprobante) y el ingreso al sistema son dos
    # pasos distintos — "recepcionado" es un estado intermedio propio, no
    # se pinta igual que "pendiente" ni igual que "completado" (que es
    # cuando además ya se confirmó el ingreso al sistema).
    if r["estado"] == "pendiente" or r.get("articulos_pendientes"):
        return "pendiente"
    if r["ingresado_sistema_en"]:
        return "completado"
    return "recepcionado"


def _tiempo_transcurrido_txt(creado_en: str) -> str:
    try:
        inicio = datetime.strptime(creado_en, "%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return "—"
    # creado_en se guarda en hora Argentina (ver _ahora_txt en database.py).
    # En Vercel el servidor corre en UTC, así que datetime.now() sin huso
    # quedaba 3 horas adelantado y este cálculo mostraba tiempos inflados.
    ahora = datetime.now(db._TZ_AR).replace(tzinfo=None)
    delta = ahora - inicio
    horas = delta.total_seconds() / 3600
    if horas < 1:
        return f"{int(delta.total_seconds() / 60)} min"
    if horas < 24:
        return f"{horas:.1f} h"
    return f"{horas / 24:.1f} d"


_DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
_MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre"]


def _hoy_ar() -> date:
    # El servidor corre en UTC (Vercel); date.today() ahí puede marcar el
    # día siguiente desde las 21hs hora Argentina en adelante.
    return datetime.now(db._TZ_AR).date()


def _grupo_fecha_txt(creado_en: str) -> str:
    try:
        fecha = datetime.strptime(creado_en, "%Y-%m-%d %H:%M:%S").date()
    except (ValueError, TypeError):
        return "—"
    hoy = _hoy_ar()
    if fecha == hoy:
        return "Hoy"
    if fecha == hoy - timedelta(days=1):
        return "Ayer"
    return f"{_DIAS[fecha.weekday()].capitalize()} {fecha.day} de {_MESES[fecha.month]}"


def _enriquecer_remito(r: dict) -> dict:
    r["nro_ingreso"] = f"ING-{r['id']:06d}"
    r["nro_remito_corto"] = r.get("numeros_remito_txt") or "—"
    r["estado_detallado"] = _estado_detallado(r)
    r["tag_clase"] = _tag_clase(r)
    r["tiempo_transcurrido_txt"] = _tiempo_transcurrido_txt(r["creado_en"])
    return r


def _rango_periodo(periodo: str, desde_param: str | None = None, hasta_param: str | None = None) -> tuple[str | None, str | None]:
    hoy = _hoy_ar()
    if periodo == "hoy":
        return hoy.isoformat(), hoy.isoformat()
    if periodo == "semana":
        return (hoy - timedelta(days=6)).isoformat(), hoy.isoformat()
    if periodo == "mes":
        return hoy.replace(day=1).isoformat(), hoy.isoformat()
    if periodo == "rango" and desde_param and hasta_param:
        return desde_param, hasta_param
    return None, None


def _archivo_valido(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXT


def _guardar_subida(file_storage, remito_id: int, prefijo: str) -> str:
    ext = Path(file_storage.filename).suffix.lower()
    nombre = f"{remito_id}_{prefijo}_{secure_filename(Path(file_storage.filename).stem)}{ext}"
    tipo_mime = file_storage.mimetype or mimetypes.guess_type(nombre)[0] or "application/octet-stream"
    db.guardar_archivo(nombre, file_storage.read(), tipo_mime)
    file_storage.stream.seek(0)
    return nombre


# --------------------------------------------------------------------- auth

@app.route("/login", methods=["GET", "POST"])
def login():
    if auth.session_valida():
        return redirect(url_for("index"))

    error = None
    timeout = request.args.get("timeout")
    next_url = request.values.get("next") or url_for("index")

    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        user = auth.verify_login(username, password)
        if user:
            auth.iniciar_sesion(user)
            return redirect(next_url)
        error = "Usuario o contraseña incorrectos."

    return render_template("login.html", error=error, timeout=timeout, next=next_url)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/cambiar-password", methods=["GET", "POST"])
@auth.login_required
def cambiar_password():
    error = None
    ok = False
    if request.method == "POST":
        actual = request.form.get("actual", "")
        nueva = request.form.get("nueva", "")
        confirmar = request.form.get("confirmar", "")
        user = auth.verify_login(session["user"], actual)
        if not user:
            error = "La contraseña actual no es correcta."
        elif len(nueva) < 6:
            error = "La nueva contraseña debe tener al menos 6 caracteres."
        elif nueva != confirmar:
            error = "Las contraseñas no coinciden."
        else:
            auth.cambiar_password(session["user"], nueva)
            ok = True

    return render_template("cambiar_password.html", error=error, ok=ok)


# --------------------------------------------------------------------- app

@app.route("/")
@auth.login_required
def index():
    estado = request.args.get("estado") or None
    q = request.args.get("q") or None
    periodo = request.args.get("periodo") or "semana"
    desde_param = request.args.get("desde") or None
    hasta_param = request.args.get("hasta") or None
    desde, hasta = _rango_periodo(periodo, desde_param, hasta_param)
    remitos = [_enriquecer_remito(r) for r in db.listar_remitos(estado=estado, q=q, desde=desde, hasta=hasta)]
    for r in remitos:
        r["grupo_fecha"] = _grupo_fecha_txt(r["creado_en"])
    return render_template(
        "index.html",
        remitos=remitos,
        resumen=db.resumen(desde=desde, hasta=hasta),
        escaneos_pendientes=db.contar_escaneos_pendientes(),
        estado=estado,
        q=q or "",
        periodo=periodo,
        desde_param=desde_param or "",
        hasta_param=hasta_param or "",
    )


@app.context_processor
def _inyectar_nav():
    # El contador de escaneos sin asignar se muestra en el nav de todas las
    # páginas (no solo el dashboard), para que se note apenas llega un mail.
    if not session.get("user"):
        return {}
    return {
        "nav_escaneos_pendientes": db.contar_escaneos_pendientes(),
        # URL del bridge de escaneo (scanner_bridge.py) — por defecto
        # localhost, pero se puede apuntar a la IP de la PC de la oficina
        # que tiene el escáner conectado, para usar el botón de escanear
        # desde otras PCs de la misma red.
        "bridge_url": os.environ.get("SCANNER_BRIDGE_URL", "http://127.0.0.1:5055"),
        "bridge_token": os.environ.get("SCANNER_API_TOKEN", ""),
    }


@app.route("/api/proveedores")
@auth.login_required
def api_proveedores():
    q = request.args.get("q", "")
    return jsonify(db.buscar_proveedores(q))


def _crear_remito_desde_archivo(
    proveedor_id: int,
    cantidad_bultos: int | None,
    contenido: bytes,
    filename: str,
    tipo_mime: str,
    usuario: str,
    tipo_control: str,
) -> int:
    """Lógica compartida entre "Nueva recepción" (sube el archivo en el
    momento) y "Asignar escaneo" (el archivo ya está guardado, llegó solo
    por mail): corre la lectura automática y crea el remito."""
    ext = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(contenido)
        tmp_path = tmp.name

    try:
        datos = ocr_parser.parse_remito(tmp_path, filename)

        numeros = [datos["nro_remito"][-5:]] if datos["nro_remito"] else []
        remito_id = db.crear_remito(
            proveedor_id=proveedor_id,
            cantidad_bultos_declarados=cantidad_bultos,
            archivo_original="",
            tipo_archivo=datos["tipo_archivo"],
            metodo_extraccion=datos["metodo_extraccion"],
            usuario_creador=usuario,
            numeros_remito=numeros,
            items=datos["items"],
            tipo_control=tipo_control,
        )

        nombre_final = f"{remito_id}_original_{secure_filename(Path(filename).stem)}{ext}"
        db.guardar_archivo(nombre_final, contenido, tipo_mime)
        db.actualizar_archivo_original(remito_id, nombre_final)
        return remito_id
    finally:
        os.unlink(tmp_path)


@app.route("/remitos/nuevo", methods=["GET", "POST"])
@auth.login_required
def nuevo_remito():
    if request.method == "GET":
        return render_template("nuevo_remito.html")

    proveedor_id = request.form.get("proveedor_id", type=int)
    cantidad_bultos = request.form.get("cantidad_bultos_declarados", type=int)
    archivo = request.files.get("archivo")
    tipo_control = request.form.get("tipo_control")

    if not proveedor_id or not db.get_proveedor(proveedor_id):
        return render_template("nuevo_remito.html", error="Elegí un proveedor válido de la lista."), 400
    if tipo_control not in ("bulto", "articulo"):
        return render_template("nuevo_remito.html", error="Elegí si vas a controlar por bulto o por artículo."), 400
    if not archivo or not archivo.filename:
        return render_template("nuevo_remito.html", error="Subí una foto o PDF del remito."), 400
    if not _archivo_valido(archivo.filename):
        return render_template("nuevo_remito.html", error="Formato de archivo no admitido."), 400

    tipo_mime = archivo.mimetype or mimetypes.guess_type(archivo.filename)[0] or "application/octet-stream"
    remito_id = _crear_remito_desde_archivo(
        proveedor_id, cantidad_bultos, archivo.read(), archivo.filename, tipo_mime, session["user"], tipo_control,
    )

    if tipo_control == "bulto":
        # Sin detalle de artículos todavía — no hay nada que revisar acá, se
        # va directo a imprimir la hoja de control para la recepción por
        # bulto. Los artículos se cargan más adelante (ver _articulos_pendientes).
        return redirect(url_for("hoja_remito", remito_id=remito_id))
    return redirect(url_for("revisar_remito", remito_id=remito_id))


@app.route("/cron/importar-escaneos")
def cron_importar_escaneos():
    # Vercel Cron llama esta ruta solo (Hobby: una vez por día, hora
    # aproximada). Si CRON_SECRET está configurado, Vercel lo manda solo en
    # sus propias llamadas — así nadie más puede disparar esto desde afuera.
    secreto = os.environ.get("CRON_SECRET")
    if secreto and request.headers.get("Authorization") != f"Bearer {secreto}":
        abort(401)
    importados = email_import.importar_escaneos_nuevos()
    return jsonify(ok=True, importados=importados)


@app.route("/api/escaneos/subir", methods=["POST"])
def api_subir_escaneo():
    # Lo llama el programa que vigila la carpeta de escaneos en la PC de la
    # oficina (watcher_escaner.py), no un usuario logueado: por eso se
    # protege con un token propio en vez de sesión.
    token = os.environ.get("SCANNER_API_TOKEN")
    if not token or request.headers.get("Authorization") != f"Bearer {token}":
        abort(401)

    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        return jsonify(ok=False, error="Falta el archivo."), 400
    if not _archivo_valido(archivo.filename):
        return jsonify(ok=False, error="Formato de archivo no admitido."), 400

    contenido = archivo.read()
    # Clave única por contenido (no por nombre de archivo): así, si el
    # programa de la PC reintenta subir el mismo escaneo por las dudas,
    # no se duplica en "Escaneos sin asignar".
    clave_unica = "scan-" + hashlib.sha256(contenido).hexdigest()
    ext = Path(archivo.filename).suffix.lower()
    nombre_guardado = f"escaneo_{clave_unica[5:21]}{ext}"
    tipo_mime = archivo.mimetype or mimetypes.guess_type(archivo.filename)[0] or "application/octet-stream"

    db.guardar_archivo(nombre_guardado, contenido, tipo_mime)
    creado = db.crear_escaneo_entrante(nombre_guardado, None, archivo.filename, clave_unica)
    return jsonify(ok=True, nuevo=bool(creado))


@app.route("/escaneos/revisar", methods=["POST"])
@auth.login_required
def escaneos_revisar():
    # Botón manual dentro de la app: no depende del cron (que en el plan
    # gratuito de Vercel corre como mucho una vez por día), así cualquiera
    # puede traer los escaneos apenas sabe que llegó un mail nuevo.
    email_import.importar_escaneos_nuevos()
    return redirect(url_for("escaneos"))


@app.route("/escaneos")
@auth.login_required
def escaneos():
    return render_template("escaneos.html", escaneos=db.listar_escaneos_pendientes())


@app.route("/escaneos/<int:escaneo_id>/asignar", methods=["GET", "POST"])
@auth.login_required
def escaneos_asignar(escaneo_id):
    escaneo = db.get_escaneo_entrante(escaneo_id)
    if not escaneo or escaneo["estado"] != "pendiente":
        abort(404)

    if request.method == "GET":
        return render_template("escaneos_asignar.html", escaneo=escaneo)

    proveedor_id = request.form.get("proveedor_id", type=int)
    cantidad_bultos = request.form.get("cantidad_bultos_declarados", type=int)
    tipo_control = request.form.get("tipo_control")
    if not proveedor_id or not db.get_proveedor(proveedor_id):
        return render_template("escaneos_asignar.html", escaneo=escaneo, error="Elegí un proveedor válido de la lista."), 400
    if tipo_control not in ("bulto", "articulo"):
        return render_template("escaneos_asignar.html", escaneo=escaneo, error="Elegí si vas a controlar por bulto o por artículo."), 400

    archivo = db.leer_archivo(escaneo["archivo_nombre"])
    if not archivo:
        return render_template("escaneos_asignar.html", escaneo=escaneo, error="No se encontró el archivo escaneado."), 404

    remito_id = _crear_remito_desde_archivo(
        proveedor_id, cantidad_bultos, archivo["contenido"], escaneo["archivo_nombre"],
        archivo["tipo_mime"], session["user"], tipo_control,
    )
    db.marcar_escaneo_asignado(escaneo_id, remito_id)

    if tipo_control == "bulto":
        return redirect(url_for("hoja_remito", remito_id=remito_id))
    return redirect(url_for("revisar_remito", remito_id=remito_id))


@app.route("/remitos/<int:remito_id>/escanear-articulos", methods=["POST"])
@auth.login_required
def escanear_articulos(remito_id):
    # Para remitos que llegaron sin detalle de artículos (o cuando el primer
    # escaneo no los leyó bien): permite escanear la factura/remito con el
    # detalle y devuelve los ítems leídos para agregar a la revisión, sin
    # crear un remito nuevo ni pisar el escaneo original.
    remito = db.get_remito(remito_id)
    if not remito:
        return jsonify(ok=False, error="Remito no encontrado."), 404

    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        return jsonify(ok=False, error="Subí una foto o PDF."), 400
    if not _archivo_valido(archivo.filename):
        return jsonify(ok=False, error="Formato de archivo no admitido."), 400

    # Se guarda como respaldo (no solo se usa para la lectura automática):
    # en los casos de control por bulto + carga posterior de artículos desde
    # otro documento (factura), este es el único registro de ese segundo
    # papel — antes se descartaba apenas se leía.
    nombre_final = _guardar_subida(archivo, remito_id, "articulos")
    db.actualizar_archivo_articulos(remito_id, nombre_final)

    ext = Path(archivo.filename).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        archivo.save(tmp.name)
        tmp_path = tmp.name
    try:
        datos = ocr_parser.parse_remito(tmp_path, archivo.filename)
    finally:
        os.unlink(tmp_path)

    return jsonify(ok=True, nro_remito=datos["nro_remito"], items=datos["items"])


@app.route("/remitos/<int:remito_id>/revisar", methods=["GET", "POST"])
@auth.login_required
def revisar_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404

    if request.method == "GET":
        return render_template(
            "revisar_remito.html", remito=remito, items=db.get_items(remito_id),
            numeros=db.listar_numeros(remito_id),
        )

    payload = request.get_json(force=True, silent=True) or {}
    items = [
        {
            "codigo_articulo": (it.get("codigo_articulo") or "").strip() or None,
            "descripcion": (it.get("descripcion") or "").strip() or None,
            "cantidad_remito": it.get("cantidad_remito"),
        }
        for it in payload.get("items", [])
    ]
    numeros = [(n or "").strip()[-5:] for n in payload.get("numeros_remito", []) if (n or "").strip()]
    db.actualizar_revision(remito_id, numeros, items)

    if remito["estado"] == "completado":
        # Remito ya cerrado (control por bulto sin detalle de artículos):
        # esto es la carga tardía de los artículos, no la revisión previa a
        # imprimir la hoja de control — no corresponde volver a esa hoja.
        # Guardar la lista todavía no cierra el control por artículo: falta
        # imprimir esa hoja, controlarla físicamente y subir la evidencia de
        # que se hizo (ver /control-articulos) — recién ahí se marca hecho.
        if remito["articulos_pendientes"] and items:
            return jsonify(ok=True, redirect=url_for("control_articulos", remito_id=remito_id))
        return jsonify(ok=True, redirect=url_for("detalle_remito", remito_id=remito_id))

    return jsonify(ok=True, redirect=url_for("hoja_remito", remito_id=remito_id))


@app.route("/remitos/<int:remito_id>/hoja")
@auth.login_required
def hoja_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    return render_template("hoja_control.html", remito=remito, items=db.get_items(remito_id))


@app.route("/remitos/<int:remito_id>/control-articulos", methods=["GET", "POST"])
@auth.login_required
def control_articulos(remito_id):
    # Último paso del control por artículo (para remitos recepcionados por
    # bulto): subir la evidencia de que la hoja de control ya impresa se
    # controló físicamente contra los artículos. No hace falta detectar
    # firma acá (a diferencia del comprobante del bulto) — alcanza con que
    # quede la foto/PDF como respaldo de que el control se hizo.
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    if remito["estado"] != "completado" or not remito["articulos_pendientes"]:
        return redirect(url_for("detalle_remito", remito_id=remito_id))
    if not db.get_items(remito_id):
        # Todavía no cargó ningún artículo — no hay nada que controlar.
        return redirect(url_for("revisar_remito", remito_id=remito_id))

    if request.method == "GET":
        return render_template("control_articulos.html", remito=remito)

    archivo = request.files.get("comprobante")
    if not archivo or not archivo.filename:
        return render_template(
            "control_articulos.html", remito=remito, error="Subí la foto o PDF de la hoja ya controlada."
        ), 400
    if not _archivo_valido(archivo.filename):
        return render_template(
            "control_articulos.html", remito=remito, error="Formato de archivo no admitido."
        ), 400

    nombre_final = _guardar_subida(archivo, remito_id, "control_articulos")
    db.marcar_articulos_completados(remito_id, nombre_final, session["user"])
    return redirect(url_for("detalle_remito", remito_id=remito_id))


@app.route("/remitos/<int:remito_id>/etiquetas")
@auth.login_required
def etiquetas_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    fecha_txt = "—"
    if remito["creado_en"]:
        f = remito["creado_en"].split(" ")[0].split("-")
        fecha_txt = f"{f[2]}/{f[1]}/{f[0]}"
    return render_template(
        "etiquetas.html",
        remito=remito,
        items=db.get_items(remito_id),
        fecha_txt=fecha_txt,
    )


@app.route("/remitos/<int:remito_id>/etiquetas.pdf")
@auth.login_required
def etiquetas_remito_pdf(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    fecha_txt = "—"
    if remito["creado_en"]:
        f = remito["creado_en"].split(" ")[0].split("-")
        fecha_txt = f"{f[2]}/{f[1]}/{f[0]}"

    pdf_bytes = generar_pdf_etiquetas(
        remito=remito,
        items=db.get_items(remito_id),
        fecha_txt=fecha_txt,
        tipo=request.args.get("tipo", "bultos"),
        ancho_mm=request.args.get("ancho", 55, type=float),
        alto_mm=request.args.get("alto", 45, type=float),
        espacio_numero=request.args.get("esp_numero", 70, type=float),
        espacio_proveedor=request.args.get("esp_proveedor", 20, type=float),
        espacio_fecha=request.args.get("esp_fecha", 10, type=float),
    )
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": "inline; filename=etiquetas.pdf"},
    )


@app.route("/remitos/<int:remito_id>/etiquetas/imprimir")
@auth.login_required
def etiquetas_remito_imprimir(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    pdf_url = url_for("etiquetas_remito_pdf", remito_id=remito_id, **request.args)
    return render_template("imprimir_etiquetas.html", pdf_url=pdf_url)


def _articulos_pendientes(remito_id: int, tipo_control: str | None) -> bool:
    # Algunos proveedores mandan el remito solo con cantidad de bultos, sin
    # detalle de artículos — ahí no hay nada que extraer/cargar en el momento
    # de la recepción. Se permite completar el control por bulto igual, pero
    # queda marcado para cargar los artículos más adelante.
    #
    # Importante: NO alcanza con mirar si el remito ya tiene ítems en la base
    # (db.get_items) para decidir esto — la lectura automática corre sobre el
    # remito original apenas se sube, antes de elegir bulto/artículo acá, y
    # puede haber detectado algo (aunque sea por error) sin que la persona
    # haya hecho ningún control por artículo real. Elegir "por bulto" deja
    # el control por artículo pendiente siempre, hasta que se complete a
    # propósito desde "Cargar artículos" (ver marcar_articulos_completados).
    return tipo_control == "bulto"


@app.route("/remitos/<int:remito_id>/confirmar", methods=["GET", "POST"])
@auth.login_required
def confirmar_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    if remito["estado"] == "completado":
        return redirect(url_for("detalle_remito", remito_id=remito_id))

    if request.method == "GET":
        return render_template("confirmar_remito.html", remito=remito, resultado=None)

    if request.form.get("manual_ok"):
        if not remito["comprobante_final"]:
            return render_template(
                "confirmar_remito.html", remito=remito,
                error="Primero subí el escaneo de la hoja completada.",
            ), 400
        observacion = (request.form.get("observacion") or "").strip()
        if not observacion:
            return render_template(
                "confirmar_remito.html", remito=remito,
                resultado={"firma_detectada": None if remito["firma_detectada"] is None else bool(remito["firma_detectada"])},
                error="Para confirmar sin detección automática, dejá una observación explicando la situación.",
            ), 400
        db.confirmar_remito(
            remito_id,
            remito["comprobante_final"],
            None if remito["firma_detectada"] is None else bool(remito["firma_detectada"]),
            confirmado_manualmente=True,
            usuario_completo=session["user"],
            observacion_confirmacion=observacion,
            articulos_pendientes=_articulos_pendientes(remito_id, remito["tipo_control"]),
        )
        return redirect(url_for("detalle_remito", remito_id=remito_id))

    # El tipo de control ya se eligió al crear el remito (no se vuelve a
    # preguntar acá) — se valida igual por si es un remito viejo, de antes de
    # que existiera ese paso.
    tipo_control = remito["tipo_control"]
    if tipo_control not in ("bulto", "articulo"):
        return render_template(
            "confirmar_remito.html", remito=remito,
            error="Este remito no tiene definido el tipo de control. Volvé al remito y cargalo de nuevo.",
        ), 400

    archivo = request.files.get("comprobante")
    if not archivo or not archivo.filename:
        return render_template(
            "confirmar_remito.html", remito=remito, error="Subí la foto o PDF de la hoja firmada."
        ), 400
    if not _archivo_valido(archivo.filename):
        return render_template(
            "confirmar_remito.html", remito=remito, error="Formato de archivo no admitido."
        ), 400

    nombre_final = _guardar_subida(archivo, remito_id, "comprobante")
    ext = Path(archivo.filename).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        archivo.save(tmp.name)
        tmp_path = tmp.name
    try:
        resultado = ocr_parser.validar_comprobante(tmp_path, archivo.filename)
    finally:
        os.unlink(tmp_path)

    if resultado["firma_detectada"]:
        db.confirmar_remito(
            remito_id,
            nombre_final,
            resultado["firma_detectada"],
            confirmado_manualmente=False,
            usuario_completo=session["user"],
            tipo_control=tipo_control,
            articulos_pendientes=_articulos_pendientes(remito_id, tipo_control),
        )
        return redirect(url_for("detalle_remito", remito_id=remito_id))

    db.guardar_intento_confirmacion(
        remito_id, nombre_final, resultado["firma_detectada"], tipo_control
    )
    remito = db.get_remito(remito_id)
    return render_template("confirmar_remito.html", remito=remito, resultado=resultado)


@app.route("/remitos/<int:remito_id>")
@auth.login_required
def detalle_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    remito = _enriquecer_remito(remito)
    return render_template("detalle_remito.html", remito=remito, items=db.get_items(remito_id))


@app.route("/remitos/<int:remito_id>/ingreso-sistema", methods=["GET", "POST"])
@auth.login_required
def ingreso_sistema(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    if remito["estado"] != "completado":
        return redirect(url_for("detalle_remito", remito_id=remito_id))
    if remito["ingresado_sistema_en"]:
        return redirect(url_for("detalle_remito", remito_id=remito_id))
    if remito["articulos_pendientes"]:
        # Remito recepcionado por bulto: falta cargar la factura con el
        # detalle de artículos y completar ese control antes de poder
        # vincularlo a una orden de compra e ingresarlo al sistema.
        return redirect(url_for("detalle_remito", remito_id=remito_id))

    if request.method == "POST":
        orden_compra_id = request.form.get("orden_compra_id", type=int)
        if not orden_compra_id or not db.get_orden_compra(orden_compra_id):
            candidatas = db.listar_ordenes_compra_por_proveedor(remito["proveedor_id"]) if remito["proveedor_id"] else []
            return render_template(
                "ingreso_sistema.html", remito=remito, candidatas=candidatas,
                total_cantidad=db.total_cantidad_items(remito_id),
                error="Elegí una orden de compra de la lista.",
            ), 400
        db.vincular_orden_compra(remito_id, orden_compra_id, session["user"], automatico=False)
        return redirect(url_for("detalle_remito", remito_id=remito_id))

    candidatas = db.listar_ordenes_compra_por_proveedor(remito["proveedor_id"]) if remito["proveedor_id"] else []
    return render_template(
        "ingreso_sistema.html", remito=remito, candidatas=candidatas,
        total_cantidad=db.total_cantidad_items(remito_id),
    )


@app.route("/remitos/<int:remito_id>/eliminar", methods=["POST"])
@auth.login_required
def eliminar_remito(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404
    if remito["estado"] == "completado" and session.get("rol") != "admin":
        return "Solo un admin puede eliminar un remito completado.", 403
    db.eliminar_remito(remito_id)
    return redirect(url_for("index"))


@app.route("/uploads/<path:filename>")
@auth.login_required
def uploads(filename):
    archivo = db.leer_archivo(filename)
    if not archivo:
        abort(404)
    return Response(archivo["contenido"], mimetype=archivo["tipo_mime"])


# ------------------------------------------------------------------- admin

@app.route("/admin/usuarios", methods=["GET", "POST"])
@auth.admin_required
def admin_usuarios():
    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        rol = request.form.get("rol", "operativo")
        if rol not in auth.ROLES:
            rol = "operativo"
        if not username or len(password) < 6:
            error = "Usuario inválido o contraseña muy corta (mínimo 6 caracteres)."
        elif db.get_usuario(username):
            error = "Ya existe un usuario con ese nombre."
        else:
            auth.crear_usuario(username, password, rol)

    return render_template("admin_usuarios.html", usuarios=db.listar_usuarios(), error=error)


@app.route("/admin/usuarios/<username>/baja", methods=["POST"])
@auth.admin_required
def admin_usuarios_baja(username):
    db.set_usuario_activo(username, False)
    return redirect(url_for("admin_usuarios"))


@app.route("/admin/usuarios/<username>/alta", methods=["POST"])
@auth.admin_required
def admin_usuarios_alta(username):
    db.set_usuario_activo(username, True)
    return redirect(url_for("admin_usuarios"))


def _parsear_proveedores_xlsx(tmp_path: str) -> list[tuple[str | None, str, str | None]]:
    import openpyxl

    wb = None
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        hoja = wb.active
        filas = []
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if not fila or len(fila) < 2:
                continue
            codigo, nombre = fila[0], fila[1]
            if not nombre or not str(nombre).strip():
                continue
            filas.append((str(codigo).strip() if codigo else None, str(nombre).strip(), None))
        return filas
    finally:
        if wb is not None:
            wb.close()


def _parsear_proveedores_html(tmp_path: str) -> list[tuple[str | None, str, str | None]]:
    # El sistema de gestión de Delmy exporta "Excel" como una tabla HTML con
    # extensión .xls (no es un binario XLS real). Se busca la fila de
    # encabezado por nombre de columna, no por posición fija.
    from bs4 import BeautifulSoup

    with open(tmp_path, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    filas_html = soup.find_all("tr")
    header_idx = None
    header = None
    for i, tr in enumerate(filas_html):
        celdas = [td.get_text(strip=True) for td in tr.find_all("td")]
        if "Código de proveedor" in celdas:
            header_idx, header = i, celdas
            break
    if header is None:
        raise ValueError('no se encontró la columna "Código de proveedor" en el archivo')

    idx_codigo = header.index("Código de proveedor")
    idx_razon = header.index("Razón social") if "Razón social" in header else None
    idx_fantasia = header.index("Nombre de fantasía") if "Nombre de fantasía" in header else None
    if idx_razon is None and idx_fantasia is None:
        raise ValueError('no se encontró "Razón social" ni "Nombre de fantasía" en el archivo')

    filas = []
    for tr in filas_html[header_idx + 1:]:
        celdas = [td.get_text(strip=True) for td in tr.find_all("td")]
        if len(celdas) <= max(idx_codigo, idx_razon or 0, idx_fantasia or 0):
            continue
        # La razón social es lo que se muestra siempre en la hoja de control
        # y en toda la app — la fantasía solo se guarda aparte para que la
        # búsqueda también encuentre al proveedor por como lo conocen.
        razon_social = (celdas[idx_razon] if idx_razon is not None else "").strip()
        nombre_fantasia = (celdas[idx_fantasia] if idx_fantasia is not None else "").strip()
        nombre = razon_social or nombre_fantasia
        if not nombre:
            continue
        codigo = celdas[idx_codigo].strip() or None
        filas.append((codigo, nombre, nombre_fantasia or None))
    return filas


_CANDIDATOS_CODIGO = ["Código", "Codigo", "Código de artículo", "Cod. Artículo", "Cod Articulo", "SKU"]
_CANDIDATOS_DESCRIPCION = ["Descripción", "Descripcion", "Nombre", "Artículo", "Articulo", "Detalle"]
_CANDIDATOS_PROVEEDOR = ["Proveedor", "Nombre de fantasía", "Razón social"]
_CANDIDATOS_CANTIDAD = ["Cantidad", "Cant.", "Cant", "Cantidad comprada", "Pedido", "Cantidad pedida"]


def _idx_header(header: list[str], candidatos: list[str]) -> int | None:
    """Compara sin importar mayúsculas/minúsculas ni espacios de más — los
    archivos reales traen "PEDIDO", "pedido ", "Pedido" todos mezclados
    según quién los tipeó."""
    header_norm = [h.strip().upper() for h in header]
    for c in candidatos:
        c_norm = c.strip().upper()
        if c_norm in header_norm:
            return header_norm.index(c_norm)
    return None


def _parsear_tabla_generica_xlsx(tmp_path: str, campos: dict[str, list[str]], obligatorios: list[str]) -> list[dict]:
    """Parser genérico reusado por artículos y por ítems de compra: busca la
    fila de encabezado entre las primeras 10 filas probando nombres de
    columna candidatos por campo (cada proveedor/sistema exporta distinto),
    y arma una lista de dicts con las claves de `campos`."""
    import openpyxl

    wb = None
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        hoja = wb.active

        header_row_idx = None
        header = None
        for i, fila in enumerate(hoja.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            if not fila:
                continue
            celdas_txt = [str(c).strip() if c is not None else "" for c in fila]
            indices = {clave: _idx_header(celdas_txt, cands) for clave, cands in campos.items()}
            if all(indices[clave] is not None for clave in obligatorios):
                header_row_idx, header = i, celdas_txt
                break
        if header is None:
            raise ValueError(
                "no se encontraron las columnas necesarias (" + ", ".join(obligatorios) + ") en el archivo"
            )

        indices = {clave: _idx_header(header, cands) for clave, cands in campos.items()}
        i_obligatorio = indices[obligatorios[0]]

        filas = []
        for fila in hoja.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not fila or i_obligatorio is None or fila[i_obligatorio] is None or not str(fila[i_obligatorio]).strip():
                continue
            registro = {}
            for clave, idx in indices.items():
                if idx is None or idx >= len(fila) or fila[idx] is None:
                    registro[clave] = None
                else:
                    registro[clave] = fila[idx]
            filas.append(registro)
        return filas
    finally:
        if wb is not None:
            wb.close()


def _parsear_tabla_generica_html(tmp_path: str, campos: dict[str, list[str]], obligatorios: list[str]) -> list[dict]:
    from bs4 import BeautifulSoup

    with open(tmp_path, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    filas_html = soup.find_all("tr")
    header_idx = None
    header = None
    for i, tr in enumerate(filas_html):
        celdas = [td.get_text(strip=True) for td in tr.find_all("td")]
        if not celdas:
            continue
        indices = {clave: _idx_header(celdas, cands) for clave, cands in campos.items()}
        if all(indices[clave] is not None for clave in obligatorios):
            header_idx, header = i, celdas
            break
    if header is None:
        raise ValueError(
            "no se encontraron las columnas necesarias (" + ", ".join(obligatorios) + ") en el archivo"
        )

    indices = {clave: _idx_header(header, cands) for clave, cands in campos.items()}
    i_obligatorio = indices[obligatorios[0]]

    filas = []
    for tr in filas_html[header_idx + 1:]:
        celdas = [td.get_text(strip=True) for td in tr.find_all("td")]
        if i_obligatorio is None or len(celdas) <= i_obligatorio or not celdas[i_obligatorio].strip():
            continue
        registro = {}
        for clave, idx in indices.items():
            registro[clave] = celdas[idx].strip() if idx is not None and idx < len(celdas) and celdas[idx].strip() else None
        filas.append(registro)
    return filas


def _es_xlsx_real(tmp_path: str) -> bool:
    with open(tmp_path, "rb") as f:
        return f.read(2) == b"PK"


_CAMPOS_COMPRA_ITEMS = {
    "codigo_articulo": _CANDIDATOS_CODIGO,
    "descripcion": _CANDIDATOS_DESCRIPCION,
    "cantidad_comprada": _CANDIDATOS_CANTIDAD,
    # SOLANO = DELMY 1, VARELA = DELMY 3 (confirmado con Nicotoli — al
    # revés de lo que decía la memoria vieja del proyecto).
    "reparto_solano": ["solano", "Solano"],
    "reparto_varela": ["varela", "Varela"],
    "reparto_deposito_col": ["deposito", "depósito", "Deposito", "Depósito"],
}


def _num_o_texto(valor) -> tuple[float | None, str | None]:
    """Intenta sacar un número limpio del valor (para poder sumar/calcular
    después); si no es un número limpio (ej. "1/2 display", "1x25", "MITAD
    CADA UNA"), no inventa nada — el número queda vacío, pero el texto
    original se guarda igual, así no se pierde el dato."""
    if valor is None:
        return None, None
    texto = str(valor).strip()
    if not texto:
        return None, None
    try:
        return float(texto.replace(",", ".")), texto
    except ValueError:
        return None, texto

_CAMPOS_CODIGOS_BARRA = {
    "codigo": _CANDIDATOS_CODIGO,
    "codigo_barra": ["Código de barras"],
}

_CAMPOS_COMBOS = {
    "combo_codigo": ["Código combo"],
    "combo_descripcion": ["Descripción combo"],
    "componente_codigo": ["Código artículo"],
    "componente_descripcion": ["Descripción artículo"],
    "cantidad": ["Cantidad"],
}


_CAMPOS_CATALOGO_REQUERIDOS = ["Código", "Proveedor", "Articulo / Detalle"]


def _num_o_none(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s == "-":
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def _parsear_catalogo_razon_social_xlsx(tmp_path: str) -> list[dict]:
    """Parser dedicado al export "...ProveedorUsoRazonSocial.xlsx": trae
    código, proveedor, código del proveedor, costos y familia/categoría —
    un código sin proveedor es un combo (regla verificada contra el archivo
    de combos: coincide 100%). Busca la fila de encabezado por nombre de
    columna en vez de posición fija, igual criterio que el resto de los
    parsers del proyecto."""
    import openpyxl

    wb = None
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        hoja = wb.active

        header_row_idx, header = None, None
        for i, fila in enumerate(hoja.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
            if not fila:
                continue
            celdas_txt = [str(c).strip() if c is not None else "" for c in fila]
            if all(c in celdas_txt for c in _CAMPOS_CATALOGO_REQUERIDOS):
                header_row_idx, header = i, celdas_txt
                break
        if header is None:
            raise ValueError(
                'no se encontraron las columnas esperadas (Código, Proveedor, Articulo / Detalle) — '
                '¿es el archivo "...ProveedorUsoRazonSocial.xlsx"?'
            )

        def idx(nombre):
            return header.index(nombre) if nombre in header else None

        i_codigo, i_proveedor, i_cod_art_prov = idx("Código"), idx("Proveedor"), idx("Cód. Art. Prov")
        i_descripcion, i_familia, i_categoria = idx("Articulo / Detalle"), idx("Familia"), idx("Categoría")
        i_costo_real, i_costo_plaza, i_precio_venta = idx("Costo Real"), idx("Costo Plaza"), idx("Precio de venta mínimo")
        i_marca = idx("Marca")

        def val(fila, i):
            return fila[i] if i is not None and i < len(fila) else None

        filas = []
        for fila in hoja.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not fila or i_codigo is None or not fila[i_codigo]:
                continue
            proveedor_txt = val(fila, i_proveedor)
            proveedor_nombre = str(proveedor_txt).strip() if proveedor_txt not in (None, "") else ""
            cod_art_prov = val(fila, i_cod_art_prov)
            filas.append({
                "codigo": str(fila[i_codigo]).strip(),
                "proveedor_nombre": proveedor_nombre,
                "cod_art_prov": str(cod_art_prov).strip() if cod_art_prov not in (None, "", "-") else None,
                "descripcion": str(val(fila, i_descripcion)).strip() if val(fila, i_descripcion) else None,
                "familia": str(val(fila, i_familia)).strip() if val(fila, i_familia) else None,
                "categoria": str(val(fila, i_categoria)).strip() if val(fila, i_categoria) else None,
                "costo_real": _num_o_none(val(fila, i_costo_real)),
                "costo_plaza": _num_o_none(val(fila, i_costo_plaza)),
                "precio_venta": _num_o_none(val(fila, i_precio_venta)),
                "marca": str(val(fila, i_marca)).strip() if val(fila, i_marca) else None,
                "es_combo": not proveedor_nombre,
            })
        return filas
    finally:
        if wb is not None:
            wb.close()


def _parsear_ordenes_compra_xlsx(tmp_path: str) -> list[dict]:
    import openpyxl

    wb = None
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        hoja = wb.active

        header_row_idx = None
        header = None
        for i, fila in enumerate(hoja.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            if fila and "# Orden de compra" in fila:
                header_row_idx, header = i, list(fila)
                break
        if header is None:
            raise ValueError('no se encontró la columna "# Orden de compra" en el archivo')

        def idx(nombre):
            return header.index(nombre) if nombre in header else None

        i_oc = idx("# Orden de compra")
        i_proveedor = idx("Nombre de fantasía")
        if i_proveedor is None:
            i_proveedor = idx("Proveedor")
        i_fecha_creacion = idx("Fecha de creación")
        i_fecha_entrega = idx("Fecha de entrega")
        i_etapa = idx("Etapa")
        i_pedido = idx("Pedido")
        i_recibido = idx("Recibido")
        i_pendiente = idx("Pendiente de recepción")
        i_importe = idx("Importe neto")

        filas = []
        for fila in hoja.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not fila or i_oc is None or not fila[i_oc]:
                continue
            filas.append(
                {
                    "nro_oc": str(fila[i_oc]).strip(),
                    "proveedor_nombre": str(fila[i_proveedor]).strip() if i_proveedor is not None and fila[i_proveedor] else "",
                    "fecha_creacion": fila[i_fecha_creacion] if i_fecha_creacion is not None else None,
                    "fecha_entrega": fila[i_fecha_entrega] if i_fecha_entrega is not None else None,
                    "etapa": fila[i_etapa] if i_etapa is not None else None,
                    "cantidad_pedida": fila[i_pedido] if i_pedido is not None else None,
                    "cantidad_recibida": fila[i_recibido] if i_recibido is not None else None,
                    "cantidad_pendiente": fila[i_pendiente] if i_pendiente is not None else None,
                    "importe_neto": fila[i_importe] if i_importe is not None else None,
                }
            )
        return filas
    finally:
        if wb is not None:
            wb.close()


@app.route("/admin/ordenes-compra", methods=["GET", "POST"])
@auth.admin_required
def admin_ordenes_compra():
    error = None
    importados = None
    vinculados = None
    if request.method == "POST":
        archivo = request.files.get("excel")
        if not archivo or not archivo.filename.lower().endswith((".xlsx", ".xlsm")):
            error = "Subí el reporte de órdenes de compra (.xlsx)."
        else:
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                archivo.save(tmp.name)
                tmp_path = tmp.name
            try:
                filas = _parsear_ordenes_compra_xlsx(tmp_path)
                importados = 0
                for fila in filas:
                    if not fila["proveedor_nombre"]:
                        continue
                    db.upsert_orden_compra(fila)
                    importados += 1
                vinculados = db.intentar_cruce_automatico()
            except Exception as e:
                error = f"No se pudo leer el archivo: {e}"
            finally:
                os.unlink(tmp_path)

    return render_template(
        "admin_ordenes_compra.html",
        ordenes=db.listar_ordenes_compra(),
        error=error,
        importados=importados,
        vinculados=vinculados,
    )


@app.route("/admin/proveedores", methods=["GET", "POST"])
@auth.admin_required
def admin_proveedores():
    error = None
    importados = None
    if request.method == "POST":
        archivo = request.files.get("excel")
        if not archivo or not archivo.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
            error = "Subí el archivo del listado de proveedores (.xlsx o .xls)."
        else:
            ext = Path(archivo.filename).suffix.lower()
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                archivo.save(tmp.name)
                tmp_path = tmp.name
            try:
                with open(tmp_path, "rb") as f:
                    es_xlsx_real = f.read(2) == b"PK"

                if es_xlsx_real:
                    filas = _parsear_proveedores_xlsx(tmp_path)
                else:
                    # Muchos sistemas de gestión exportan una tabla HTML con
                    # extensión .xls en vez de un binario real.
                    filas = _parsear_proveedores_html(tmp_path)

                importados = 0
                for codigo, nombre, nombre_fantasia in filas:
                    db.upsert_proveedor(codigo, nombre, nombre_fantasia)
                    importados += 1
            except Exception as e:
                error = f"No se pudo leer el archivo: {e}"
            finally:
                os.unlink(tmp_path)

    return render_template(
        "admin_proveedores.html",
        proveedores=db.listar_proveedores(),
        error=error,
        importados=importados,
    )


# --------------------------------------------------------- laboratorio (admin)
#
# Zona de trabajo exclusiva del admin para probar y desarrollar
# funcionalidades nuevas antes de llevarlas al flujo operativo normal. No
# reemplaza ni interfiere con las pantallas de recepción existentes.

ARCHIVOS_COMPRA_EXT = ALLOWED_EXT | {".xlsx", ".xlsm", ".xls"}


def _archivo_compra_valido(filename: str) -> bool:
    return Path(filename).suffix.lower() in ARCHIVOS_COMPRA_EXT


_UNIDADES_MEDIDA_TXT = r"(?:cm|mm|ml|kg|gr|g|lt|l|cc|m)"
_PATRON_MULTIPLICADOR = re.compile(
    rf"(?:(?<!\d)x\s*(\d+)\b(?!\s*{_UNIDADES_MEDIDA_TXT}\b)|"
    rf"caja\s*(?:de)?\s*(\d+)(?!\s*{_UNIDADES_MEDIDA_TXT}\b)|"
    rf"pack\s*(?:de)?\s*x?\s*(\d+)(?!\s*{_UNIDADES_MEDIDA_TXT}\b)|"
    rf"bl[ií]ster\s*(?:de)?\s*x?\s*(\d+)(?!\s*{_UNIDADES_MEDIDA_TXT}\b)|"
    rf"bolsa\s*(?:de)?\s*x?\s*(\d+)(?!\s*{_UNIDADES_MEDIDA_TXT}\b))",
    re.IGNORECASE,
)


def _detectar_multiplicador(texto: str | None) -> int | None:
    """Busca en la descripción TAL COMO viene en el remito un patrón de
    empaque (X10, caja de 50, pack x 12, blíster x 50, bolsa de 100...).
    Es solo un dato a la vista para que la persona decida — nunca se usa
    para multiplicar la cantidad sola, eso sería inventar un dato."""
    if not texto:
        return None
    m = _PATRON_MULTIPLICADOR.search(texto)
    if not m:
        return None
    for g in m.groups():
        if g:
            return int(g)
    return None


# --------------------------------------------- clasificación de descripción
#
# Separa una descripción en capas (artículo, diseño, atributo, medida,
# código) para poder: 1) buscar códigos que vienen escondidos DENTRO de la
# descripción en vez de en el campo código, y 2) comparar dos descripciones
# dándole más peso a lo que realmente identifica el artículo (el tipo de
# producto y su código) que a atributos secundarios (color, presentación).
# Las listas de palabras salen del vocabulario real del catálogo de Delmy,
# no son genéricas — van a necesitar ajuste con el tiempo si aparece
# vocabulario nuevo.

_STOPWORDS_DESC = {"CON", "PARA", "DE", "LA", "EL", "LOS", "LAS", "Y", "EN", "POR", "DEL", "AL", "UN", "UNA", "SIN"}
_COLORES = {
    "ROJO", "AZUL", "VERDE", "AMARILLO", "NEGRO", "BLANCO", "ROSA", "FUCSIA", "VIOLETA", "NARANJA",
    "GRIS", "MARRON", "BORDO", "TURQUESA", "MULTICOLOR", "TRANSPARENTE", "DORADO", "PLATEADO",
    "PLATA", "METALIZADO", "PERLADO", "FLUO", "PASTEL", "CELESTE", "CHOCOLATE", "LILA", "BEIGE",
    "CORAL", "MOSTAZA", "OLIVA", "INDIGO", "MAGENTA", "CIAN", "ESMERALDA", "CREMA", "HUESO",
    "KAKI", "COBRE", "BRONCE",
}
_MATERIALES = {
    "PAPEL", "METAL", "GOMA", "PLASTICO", "PLASTICA", "GLITTER", "TELA", "MADERA", "CERAMICA",
    "VIDRIO", "LATEX", "FOAM", "CARTON", "CARTULINA",
}
_PRESENTACION = {
    "SET", "CAJA", "BOLSA", "BLISTER", "PACK", "KIT", "UNIDAD", "UNIDADES", "UNID", "PIEZAS",
    "GRANDE", "MINI", "CHICO", "CHICA", "MEDIANO", "MEDIANA",
}
_ATRIBUTOS_DESC = _COLORES | _MATERIALES | _PRESENTACION
_UNIDADES_MEDIDA = {"CM", "MM", "ML", "KG", "GR", "GRS", "GRAMOS", "GRAMO", "G", "LT", "L", "CC", "M"}
_PATRON_MEDIDA = re.compile(r"^T\d+$")
# Medidas tipo "10x15" o "10X15" (ancho x alto, en un solo token) — se
# revisan ANTES que el patrón de código, porque si no "10X15" se clasifica
# como código (tiene dígitos y letras) en vez de como medida.
_PATRON_DIMENSION = re.compile(r"^\d{1,4}X\d{1,4}$")
_PATRON_CODIGO_TOKEN = re.compile(r"^(?=.*\d)[A-Z0-9]{3,}$")


def _es_atributo(tok: str) -> bool:
    if tok in _ATRIBUTOS_DESC:
        return True
    if tok.endswith("S") and tok[:-1] in _ATRIBUTOS_DESC:  # plurales simples (DORADAS -> DORADO)
        return True
    return False


def _clasificar_descripcion(texto: str | None) -> dict:
    resultado = {"articulo": None, "disenio": [], "atributo": [], "medida": [], "codigo": []}
    if not texto:
        return resultado
    tokens = re.findall(r"[A-ZÁÉÍÓÚÑ0-9]+", texto.upper())
    i, n = 0, len(tokens)
    while i < n:
        tok = tokens[i]
        siguiente = tokens[i + 1] if i + 1 < n else None
        siguiente2 = tokens[i + 2] if i + 2 < n else None
        if tok in _STOPWORDS_DESC:
            i += 1
            continue
        # Medida tipo "10 X 15" escrita con espacios (3 tokens separados).
        if re.fullmatch(r"\d{1,4}", tok) and siguiente == "X" and siguiente2 and re.fullmatch(r"\d{1,4}", siguiente2):
            resultado["medida"].append(f"{tok}X{siguiente2}")
            i += 3
            continue
        # Medida tipo "10x15" pegada, en un solo token.
        if _PATRON_DIMENSION.match(tok):
            resultado["medida"].append(tok)
            i += 1
            continue
        if re.fullmatch(r"\d+", tok) and siguiente in _UNIDADES_MEDIDA:
            resultado["medida"].append(tok + siguiente)
            i += 2
            continue
        if _PATRON_MEDIDA.match(tok):
            resultado["medida"].append(tok)
            i += 1
            continue
        if _es_atributo(tok):
            resultado["atributo"].append(tok)
            i += 1
            continue
        if _PATRON_CODIGO_TOKEN.match(tok):
            resultado["codigo"].append(tok)
            i += 1
            continue
        if resultado["articulo"] is None and tok.isalpha() and len(tok) >= 3:
            resultado["articulo"] = tok
            i += 1
            continue
        resultado["disenio"].append(tok)
        i += 1
    return resultado


def _similitud_por_capas(desc_a: str | None, desc_b: str | None) -> float:
    """Similitud entre dos descripciones, dándole el peso más grande a lo
    que realmente distingue UNA variante de otra dentro de la misma línea
    de producto:

    - Medida (10x15 vs 14x20): si las dos tienen medida y NO coincide, son
      variantes distintas seguro — resta fuerte.
    - Diseño (el nombre propio: personaje, estampa, modelo — ej. "Snoopy"
      vs "Boho"): si no comparten NINGUNA palabra, es casi seguro un
      artículo distinto aunque el resto del texto (marca, tipo, medida) se
      parezca mucho — resta fuerte. Esto es lo que evita recomendar
      "Agenda Mood" cuando el remito dice "Agenda Snoopy".

    El texto plano ya no es lo que más pesa — es apenas un piso, porque dos
    descripciones de variantes distintas del mismo producto comparten casi
    todo el texto salvo la palabra que las distingue."""
    base = SequenceMatcher(None, (desc_a or "").strip().upper(), (desc_b or "").strip().upper()).ratio()
    if not desc_a or not desc_b:
        return base
    ca, cb = _clasificar_descripcion(desc_a), _clasificar_descripcion(desc_b)
    puntaje = base * 0.4

    if ca["articulo"] and ca["articulo"] == cb["articulo"]:
        puntaje += 0.20
    elif ca["articulo"] and cb["articulo"]:
        puntaje -= 0.15  # tipos de producto distintos (ej. "Agenda" vs "Cuaderno")

    if ca["medida"] and cb["medida"]:
        if set(ca["medida"]) & set(cb["medida"]):
            puntaje += 0.15
        else:
            puntaje -= 0.30

    disenio_a, disenio_b = set(ca["disenio"]), set(cb["disenio"])
    if disenio_a and disenio_b:
        comunes = disenio_a & disenio_b
        if comunes:
            puntaje += 0.20 * (len(comunes) / max(len(disenio_a), len(disenio_b)))
        else:
            puntaje -= 0.35

    atributo_a, atributo_b = set(ca["atributo"]), set(cb["atributo"])
    if atributo_a and atributo_b:
        puntaje += 0.05 if (atributo_a & atributo_b) else -0.10

    if set(ca["codigo"]) & set(cb["codigo"]):
        puntaje += 0.10
    return max(0.0, min(1.0, puntaje))


def _matchear_articulo(codigo: str | None, descripcion: str | None, proveedor_id: int | None) -> dict:
    """Busca si un ítem leído de un remito/factura existe en el catálogo de
    artículos de Delmy, en capas — nunca inventa ni elige a ciegas:

    1) Código de PROVEEDOR exacto (lo que trae el remito suele ser el
       código que usa el proveedor, no el interno de Delmy).
    2) Código interno de Delmy exacto (por si el remito ya trae ese).
    3) Código escondido DENTRO de la descripción (a veces el código real no
       está en el campo código sino como una sigla/número dentro del texto,
       ej. "MARCADOR SHARPIE... N 1812793" — el clasificador de descripción
       lo separa y se prueba igual que un código de proveedor).
    4) Código parcial — contención en ambos sentidos, para cuando el
       código del remito viene incompleto o con caracteres de más.
    5) Recomendación por descripción — varios candidatos, comparando por
       capas (más peso al tipo de artículo y a los códigos compartidos que
       al color/material), para artículos nuevos o con código viejo que ya
       no está en el catálogo.

    Si lo encontrado es un combo, adjunta de qué componentes se arma. Si la
    descripción del remito trae un patrón de empaque (X10, caja de 50...),
    lo señala como dato a la vista, sin aplicarlo solo."""
    resultado = {
        "estado": "no_encontrado",
        "catalogo_codigo": None,
        "catalogo_cod_art_prov": None,
        "catalogo_descripcion": None,
        "similitud": None,
        "candidatos": [],
        "combo_componentes": [],
        "multiplicador_detectado": _detectar_multiplicador(descripcion),
    }
    codigo_norm = (codigo or "").strip()

    def _con_combo(art_codigo: str):
        comp = db.get_combo_componentes(art_codigo)
        if comp:
            resultado["combo_componentes"] = comp

    if codigo_norm and proveedor_id:
        art = db.buscar_por_cod_art_prov(codigo_norm, proveedor_id)
        if art:
            resultado["catalogo_codigo"] = art["codigo"]
            resultado["catalogo_cod_art_prov"] = art.get("cod_art_prov")
            resultado["catalogo_descripcion"] = art["descripcion"]
            resultado["estado"] = "codigo_proveedor_exacto"
            _con_combo(art["codigo"])
            return resultado

    if codigo_norm:
        art = db.buscar_articulo_exacto(codigo_norm, proveedor_id)
        if art:
            resultado["catalogo_codigo"] = art["codigo"]
            resultado["catalogo_cod_art_prov"] = art.get("cod_art_prov")
            resultado["catalogo_descripcion"] = art["descripcion"]
            if descripcion and art["descripcion"]:
                ratio = _similitud_por_capas(descripcion, art["descripcion"])
                resultado["similitud"] = round(ratio, 2)
                resultado["estado"] = "codigo_exacto" if ratio >= 0.5 else "codigo_exacto_descripcion_distinta"
            else:
                resultado["estado"] = "codigo_exacto"
            _con_combo(art["codigo"])
            return resultado

    if descripcion and proveedor_id:
        codigos_en_desc = _clasificar_descripcion(descripcion)["codigo"]
        for cod_embebido in codigos_en_desc:
            art = db.buscar_por_cod_art_prov(cod_embebido, proveedor_id) or db.buscar_articulo_exacto(cod_embebido, proveedor_id)
            if art:
                # Un código suelto (ej. "500") puede coincidir de pura
                # casualidad con el código de otro artículo del mismo
                # proveedor que no tiene nada que ver — si la descripción
                # no se parece en nada, no se da por bueno a ciegas.
                ratio = _similitud_por_capas(descripcion, art["descripcion"]) if art["descripcion"] else 0.0
                resultado["catalogo_codigo"] = art["codigo"]
                resultado["catalogo_cod_art_prov"] = art.get("cod_art_prov")
                resultado["catalogo_descripcion"] = art["descripcion"]
                resultado["similitud"] = round(ratio, 2)
                resultado["estado"] = "codigo_en_descripcion" if ratio >= 0.35 else "codigo_en_descripcion_dudoso"
                _con_combo(art["codigo"])
                return resultado

    if codigo_norm and proveedor_id and len(codigo_norm) >= 3:
        candidatos = db.buscar_cod_art_prov_parcial(codigo_norm, proveedor_id)
        if candidatos:
            resultado["candidatos"] = candidatos
            resultado["estado"] = "codigo_parcial"
            return resultado

    if descripcion and proveedor_id:
        articulos_proveedor = db.listar_articulos_por_proveedor(proveedor_id)
        puntuados = []
        for art in articulos_proveedor:
            if not art.get("descripcion"):
                continue
            ratio = _similitud_por_capas(descripcion, art["descripcion"])
            puntuados.append({
                "codigo": art["codigo"], "cod_art_prov": art.get("cod_art_prov"),
                "descripcion": art["descripcion"], "similitud": round(ratio, 2),
            })
        puntuados.sort(key=lambda x: x["similitud"], reverse=True)
        candidatos = puntuados[:5]
        if candidatos:
            resultado["candidatos"] = candidatos
            mejor = candidatos[0]
            if mejor["similitud"] >= 0.6:
                resultado["catalogo_codigo"] = mejor["codigo"]
                resultado["catalogo_cod_art_prov"] = mejor.get("cod_art_prov")
                resultado["catalogo_descripcion"] = mejor["descripcion"]
                resultado["similitud"] = mejor["similitud"]
                resultado["estado"] = "descripcion_similar"
            else:
                resultado["estado"] = "recomendacion_baja_confianza"

    return resultado


def _obtener_archivo_subido(campo_archivo: str = "excel") -> tuple[str | None, bytes | None, str | None]:
    """Devuelve (nombre_archivo, contenido, error). Soporta tanto un
    archivo subido directo (formularios chicos, sigue andando igual que
    siempre) como uno reconstruido a partir de trozos (subida_id) — para
    esquivar el límite de 4.5MB por request que tiene cualquier función
    serverless de Vercel."""
    subida_id = request.form.get("subida_id")
    if subida_id:
        nombre, contenido = db.reconstruir_archivo(subida_id)
        db.borrar_chunks(subida_id)
        if not contenido:
            return None, None, "No se recibió ningún trozo del archivo — probá subirlo de nuevo."
        return nombre, contenido, None

    archivo = request.files.get(campo_archivo)
    if not archivo or not archivo.filename:
        return None, None, "Subí un archivo."
    return archivo.filename, archivo.read(), None


@app.route("/admin/laboratorio/subir-chunk", methods=["POST"])
@auth.admin_required
def admin_laboratorio_subir_chunk():
    subida_id = request.form.get("subida_id")
    indice = request.form.get("indice", type=int)
    total = request.form.get("total", type=int)
    nombre_archivo = request.form.get("nombre_archivo")
    chunk = request.files.get("chunk")

    if not subida_id or indice is None or not total or not chunk:
        return jsonify(ok=False, error="Falta un dato del trozo."), 400

    db.guardar_chunk(subida_id, indice, total, nombre_archivo, chunk.read())
    recibidos, total_guardado = db.contar_chunks(subida_id)
    return jsonify(ok=True, recibidos=recibidos, total=total_guardado)


@app.route("/admin/laboratorio")
@auth.admin_required
def admin_laboratorio():
    return render_template(
        "admin_laboratorio.html",
        cantidad_articulos=db.contar_articulos(),
        cantidad_inconsistencias=db.contar_inconsistencias_abiertas(),
        cantidad_propuestas=db.contar_propuestas_pendientes(),
    )


@app.route("/admin/laboratorio/articulos", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_articulos():
    # Importa el catálogo real de Delmy (export "...ProveedorUsoRazonSocial.
    # xlsx" del sistema de gestión) — código, proveedor, código del
    # proveedor, costos y familia/categoría. Un código sin proveedor es un
    # combo. Cada importación queda registrada como una "carga" con su
    # historial de cambios campo por campo e inconsistencias detectadas
    # (proveedor no reconocido, combo con proveedor, sin descripción, costo
    # mayor al precio de venta, cambios de costo bruscos) — nada se pierde
    # ni se sobrescribe en silencio.
    error = None
    resultado = None
    if request.method == "POST":
        nombre_archivo, contenido, error = _obtener_archivo_subido("excel")
        if not error and not nombre_archivo.lower().endswith((".xlsx", ".xlsm")):
            error = "Subí el archivo del catálogo (\"...ProveedorUsoRazonSocial.xlsx\")."
        if not error:
            ext = Path(nombre_archivo).suffix.lower()
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(contenido)
                tmp_path = tmp.name
            try:
                filas = _parsear_catalogo_razon_social_xlsx(tmp_path)
                if not filas:
                    error = "El archivo no tiene ninguna fila con código — revisá que sea el correcto."
                else:
                    carga_id = db.crear_carga_articulos(nombre_archivo, session["user"])
                    resultado = db.procesar_carga_articulos(carga_id, filas)
                    resultado["carga_id"] = carga_id
            except Exception as e:
                error = f"No se pudo leer el archivo: {e}"
            finally:
                os.unlink(tmp_path)

    return render_template(
        "admin_laboratorio_articulos.html",
        articulos=db.listar_articulos(limite=100),
        cantidad_articulos=db.contar_articulos(),
        cantidad_inconsistencias=db.contar_inconsistencias_abiertas(),
        error=error,
        resultado=resultado,
    )


@app.route("/admin/laboratorio/articulos/cargas")
@auth.admin_required
def admin_laboratorio_articulos_cargas():
    return render_template("admin_laboratorio_articulos_cargas.html", cargas=db.listar_articulos_cargas())


@app.route("/admin/laboratorio/articulos/inconsistencias", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_articulos_inconsistencias():
    if request.method == "POST":
        inconsistencia_id = request.form.get("inconsistencia_id", type=int)
        if inconsistencia_id:
            db.marcar_inconsistencia_resuelta(inconsistencia_id)
        return redirect(url_for("admin_laboratorio_articulos_inconsistencias"))

    ver_todas = request.args.get("todas") == "1"
    return render_template(
        "admin_laboratorio_articulos_inconsistencias.html",
        inconsistencias=db.listar_inconsistencias(solo_abiertas=not ver_todas),
        ver_todas=ver_todas,
        cantidad_abiertas=db.contar_inconsistencias_abiertas(),
    )


# ------------------------------------------------------------- propuestas
#
# Cola de propuestas de alta/corrección generadas desde Identificación e
# inventario y Revisión de remitos cuando un artículo no se encuentra. Se
# exportan en el formato EXACTO de "Importación Masiva" (39 columnas) —
# solo se completan los campos que ya tenemos, el resto queda vacío para
# terminarlo en el sistema de gestión real. Nunca se inventa un dato en un
# campo que no tenemos.

_COLUMNAS_IMPORTACION_MASIVA = [
    "Código interno", "Tipo", "Familia", "Categoría", "Marca", "Sinónimo", "Descripción",
    "Descripción fiscal", "Meses de garantía", "Controlar stock", "No permitir venderlo",
    "No permitir comprarlo", "Es fraccionable", "Detalles", "Magnitud", "Unidad de Medida",
    "Peso Neto(kg)", "Envoltorio", "Peso Bruto(kg)", "Ancho (mm)", "Alto (mm)", "Profundidad (mm)",
    "Proveedor", "Cod. Art. Prov", "Moneda costo", "Lista proveedor neto", "% Descuento 1",
    "% Descuento 2", "% Descuento 3", "% Descuento 4", "% Ajuste + / -", "% utilidad",
    "Disponibilidad (stock)", "Observaciones del proveedor", "01 - MOSTRADOR", "Alicuota de IVA",
    "IVA por actividad", "Fecha de creación", "Fecha de modificación",
]


def _generar_excel_importacion_masiva(propuestas: list[dict]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Importación Masiva"
    ws.append(_COLUMNAS_IMPORTACION_MASIVA)
    for p in propuestas:
        fila = {col: None for col in _COLUMNAS_IMPORTACION_MASIVA}
        fila["Código interno"] = p["codigo"] if p["tipo"] == "modificacion" else None
        fila["Tipo"] = "Bienes"
        fila["Familia"] = p.get("familia")
        fila["Categoría"] = p.get("categoria")
        fila["Marca"] = p.get("marca")
        fila["Descripción"] = p.get("descripcion")
        fila["Proveedor"] = p.get("proveedor_nombre")
        fila["Cod. Art. Prov"] = p.get("cod_art_prov")
        fila["Lista proveedor neto"] = p.get("costo")
        fila["01 - MOSTRADOR"] = p.get("precio_venta")
        fila["Observaciones del proveedor"] = p.get("observacion")
        ws.append([fila[col] for col in _COLUMNAS_IMPORTACION_MASIVA])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.route("/admin/laboratorio/propuestas", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_propuestas():
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "eliminar":
            db.eliminar_propuesta(request.form.get("propuesta_id", type=int))
        return redirect(url_for("admin_laboratorio_propuestas"))

    ver_todas = request.args.get("todas") == "1"
    return render_template(
        "admin_laboratorio_propuestas.html",
        propuestas=db.listar_propuestas(solo_pendientes=not ver_todas),
        ver_todas=ver_todas,
        cantidad_pendientes=db.contar_propuestas_pendientes(),
    )


@app.route("/admin/laboratorio/propuestas/crear", methods=["POST"])
@auth.admin_required
def admin_laboratorio_propuestas_crear():
    datos = request.get_json(silent=True) or {}
    if not datos.get("descripcion") and not datos.get("codigo"):
        return jsonify(ok=False, error="Falta al menos el código o la descripción."), 400

    propuesta_id = db.crear_propuesta_articulo(
        tipo=datos.get("tipo") or "nuevo",
        codigo=datos.get("codigo") or None,
        descripcion=datos.get("descripcion") or None,
        proveedor_nombre=datos.get("proveedor_nombre") or None,
        cod_art_prov=datos.get("cod_art_prov") or None,
        familia=datos.get("familia") or None,
        categoria=datos.get("categoria") or None,
        marca=datos.get("marca") or None,
        costo=float(datos["costo"]) if datos.get("costo") not in (None, "") else None,
        precio_venta=float(datos["precio_venta"]) if datos.get("precio_venta") not in (None, "") else None,
        observacion=datos.get("observacion") or None,
        origen=datos.get("origen") or "manual",
        origen_referencia=datos.get("origen_referencia") or None,
        usuario=session["user"],
    )
    return jsonify(ok=True, id=propuesta_id)


@app.route("/admin/laboratorio/propuestas/exportar")
@auth.admin_required
def admin_laboratorio_propuestas_exportar():
    ids_param = request.args.get("ids")
    if ids_param:
        propuestas = db.get_propuestas_por_ids([int(i) for i in ids_param.split(",") if i.strip().isdigit()])
    else:
        propuestas = db.listar_propuestas(solo_pendientes=True)

    if not propuestas:
        return "No hay propuestas para exportar", 400

    contenido = _generar_excel_importacion_masiva(propuestas)
    db.marcar_propuestas_exportadas([p["id"] for p in propuestas])
    return Response(
        contenido,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=propuestas_importacion_masiva_{date.today().isoformat()}.xlsx"},
    )


@app.route("/admin/laboratorio/articulos/codigos-barra", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_articulos_codigos_barra():
    # Importa el archivo "exportacion_de_articulos.xlsx" (el de Marca, no el
    # de RazonSocial) — trae varias filas por artículo, una por cada código
    # de barras/código alternativo que tiene. Se usa para identificar el
    # artículo al escanear en el módulo de inventario.
    error = None
    importados = None
    if request.method == "POST":
        nombre_archivo, contenido, error = _obtener_archivo_subido("excel")
        if not error and not nombre_archivo.lower().endswith((".xlsx", ".xlsm")):
            error = 'Subí el archivo del catálogo con códigos de barras (el que tiene la columna "Marca").'
        if not error:
            ext = Path(nombre_archivo).suffix.lower()
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(contenido)
                tmp_path = tmp.name
            try:
                if _es_xlsx_real(tmp_path):
                    filas = _parsear_tabla_generica_xlsx(tmp_path, _CAMPOS_CODIGOS_BARRA, obligatorios=["codigo", "codigo_barra"])
                else:
                    filas = _parsear_tabla_generica_html(tmp_path, _CAMPOS_CODIGOS_BARRA, obligatorios=["codigo", "codigo_barra"])
                importados = db.importar_codigos_barra(filas)
            except Exception as e:
                error = f"No se pudo leer el archivo: {e}"
            finally:
                os.unlink(tmp_path)

    return render_template(
        "admin_laboratorio_codigos_barra.html",
        cantidad_codigos_barra=db.contar_codigos_barra(),
        error=error,
        importados=importados,
    )


@app.route("/admin/laboratorio/articulos/combos", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_articulos_combos():
    # Importa el archivo "exportacion_de_combos_y_componentes.xlsx" — trae
    # una fila de cabecera "Combo" (sin componente) y una o más filas
    # "Artículo" con el componente real y la cantidad. Solo se guardan las
    # de tipo "Artículo" (se descartan las de cabecera al importar).
    error = None
    importados = None
    if request.method == "POST":
        nombre_archivo, contenido, error = _obtener_archivo_subido("excel")
        if not error and not nombre_archivo.lower().endswith((".xlsx", ".xlsm")):
            error = "Subí el archivo de combos y componentes (.xlsx)."
        if not error:
            ext = Path(nombre_archivo).suffix.lower()
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(contenido)
                tmp_path = tmp.name
            try:
                if _es_xlsx_real(tmp_path):
                    filas = _parsear_tabla_generica_xlsx(tmp_path, _CAMPOS_COMBOS, obligatorios=["combo_codigo", "componente_codigo"])
                else:
                    filas = _parsear_tabla_generica_html(tmp_path, _CAMPOS_COMBOS, obligatorios=["combo_codigo", "componente_codigo"])
                importados = db.importar_combos_componentes(filas)
            except Exception as e:
                error = f"No se pudo leer el archivo: {e}"
            finally:
                os.unlink(tmp_path)

    return render_template(
        "admin_laboratorio_combos.html",
        cantidad_combos=db.contar_combos_componentes(),
        error=error,
        importados=importados,
    )


@app.route("/admin/laboratorio/inventario")
@auth.admin_required
def admin_laboratorio_inventario():
    return render_template(
        "admin_laboratorio_inventario.html",
        cantidad_articulos=db.contar_articulos(),
        cantidad_codigos_barra=db.contar_codigos_barra(),
    )


# ---------------------------------------- revisión de remitos reales (admin)
#
# Toma los remitos que Javier ya escaneó de verdad en la recepción (no una
# subida de prueba) y los pasa por el mismo motor de matching contra el
# catálogo que usa el sandbox de "Reconocimiento", más el estado real del
# cruce con la orden de compra — para poder auditar si lo que se detectó,
# el catálogo, y la OC, coinciden.

@app.route("/admin/laboratorio/revision-remitos")
@auth.admin_required
def admin_laboratorio_revision_remitos():
    hasta = request.args.get("hasta") or date.today().isoformat()
    desde = request.args.get("desde") or (date.today() - timedelta(days=30)).isoformat()
    q = request.args.get("q", "")
    remitos = db.listar_remitos(q=q or None, desde=desde, hasta=hasta)
    return render_template(
        "admin_laboratorio_revision_remitos.html",
        remitos=remitos,
        desde=desde,
        hasta=hasta,
        q=q,
    )


@app.route("/admin/laboratorio/revision-remitos/<int:remito_id>")
@auth.admin_required
def admin_laboratorio_revision_remito_detalle(remito_id):
    remito = db.get_remito(remito_id)
    if not remito:
        return "Remito no encontrado", 404

    confirmaciones = db.get_confirmaciones_remito(remito_id)
    items = []
    for it in db.get_items(remito_id):
        match = _matchear_articulo(it.get("codigo_articulo"), it.get("descripcion"), remito["proveedor_id"])
        confirmacion = confirmaciones.get(it["nro_orden"])
        items.append({**it, **match, "confirmado": confirmacion})

    return render_template(
        "admin_laboratorio_revision_remito_detalle.html",
        remito=remito,
        items=items,
    )


@app.route("/admin/laboratorio/revision-remitos/<int:remito_id>/confirmar", methods=["POST"])
@auth.admin_required
def admin_laboratorio_revision_confirmar(remito_id):
    datos = request.get_json(silent=True) or {}
    nro_orden = datos.get("nro_orden")
    codigo = (datos.get("codigo") or "").strip()
    if nro_orden is None or not codigo:
        return jsonify(ok=False, error="Falta el ítem o el código a confirmar."), 400
    db.confirmar_match_revision(remito_id, int(nro_orden), codigo, session["user"])
    return jsonify(ok=True)


@app.route("/admin/laboratorio/revision-remitos/<int:remito_id>/quitar-confirmacion", methods=["POST"])
@auth.admin_required
def admin_laboratorio_revision_quitar_confirmacion(remito_id):
    datos = request.get_json(silent=True) or {}
    nro_orden = datos.get("nro_orden")
    if nro_orden is None:
        return jsonify(ok=False, error="Falta el ítem."), 400
    db.quitar_confirmacion_revision(remito_id, int(nro_orden))
    return jsonify(ok=True)


@app.route("/admin/laboratorio/revision-remitos/calendario")
@auth.admin_required
def admin_laboratorio_revision_calendario():
    mes_param = request.args.get("mes")
    if mes_param:
        anio, mes = (int(x) for x in mes_param.split("-"))
    else:
        hoy = date.today()
        anio, mes = hoy.year, hoy.month

    primer_dia = date(anio, mes, 1)
    ultimo_dia_num = calendar.monthrange(anio, mes)[1]
    ultimo_dia = date(anio, mes, ultimo_dia_num)

    agrupado = db.remitos_agrupados_por_dia(primer_dia.isoformat(), ultimo_dia.isoformat())

    cal = calendar.Calendar(firstweekday=0)  # lunes primero
    semanas = cal.monthdayscalendar(anio, mes)  # 0 = día fuera del mes

    mes_anterior = (date(anio, mes, 1) - timedelta(days=1))
    mes_siguiente = (ultimo_dia + timedelta(days=1))

    return render_template(
        "admin_laboratorio_revision_calendario.html",
        anio=anio, mes=mes,
        nombre_mes=["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
                    "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"][mes - 1],
        semanas=semanas,
        agrupado=agrupado,
        mes_anterior=f"{mes_anterior.year}-{mes_anterior.month:02d}",
        mes_siguiente=f"{mes_siguiente.year}-{mes_siguiente.month:02d}",
        hoy=date.today().isoformat(),
    )


@app.route("/api/ubicaciones/buscar")
@auth.admin_required
def api_ubicaciones_buscar():
    q = request.args.get("q", "")
    return jsonify(db.buscar_ubicaciones(q))


@app.route("/api/articulos/identificar")
@auth.admin_required
def api_articulos_identificar():
    q = request.args.get("q", "")
    return jsonify(db.buscar_articulo_identificacion(q))


@app.route("/admin/laboratorio/inventario/guardar", methods=["POST"])
@auth.admin_required
def admin_laboratorio_inventario_guardar():
    datos = request.get_json(silent=True) or {}
    ubicacion_id = datos.get("ubicacion_id")
    codigo_articulo = (datos.get("codigo_articulo") or "").strip()
    cantidad = datos.get("cantidad")
    posicion = (datos.get("posicion") or "").strip()

    if not ubicacion_id or not codigo_articulo:
        return jsonify(ok=False, error="Falta la ubicación o el artículo."), 400
    try:
        cantidad = float(cantidad)
    except (TypeError, ValueError):
        return jsonify(ok=False, error="La cantidad no es válida."), 400

    # Si se cargó una posición más fina (ej. escaneaste 01-30-10 pero el
    # artículo va en la posición 20 adentro de ese piso), se resuelve (o se
    # crea sola, si todavía no estaba cargada como etiqueta) antes de
    # guardar — así no hace falta pasar antes por la pantalla de
    # Ubicaciones para poder inventariar en el momento.
    ubicacion_final = ubicacion_id
    ubicacion_codigo_final = None
    if posicion:
        sububicacion, error = db.obtener_o_crear_sububicacion(ubicacion_id, posicion)
        if error:
            return jsonify(ok=False, error=error), 400
        ubicacion_final = sububicacion["id"]
        ubicacion_codigo_final = sububicacion["codigo"]

    registro_id = db.crear_registro_inventario(ubicacion_final, codigo_articulo, cantidad, session["user"])
    return jsonify(ok=True, id=registro_id, ubicacion_codigo=ubicacion_codigo_final)


@app.route("/admin/laboratorio/inventario/exportar")
@auth.admin_required
def admin_laboratorio_inventario_exportar():
    hoy = date.today().isoformat()
    desde = request.args.get("desde") or hoy
    hasta = request.args.get("hasta") or hoy
    filas = db.listar_inventario_para_exportar(desde, hasta)
    if not filas:
        return "No hay nada inventariado en ese rango de fechas.", 400

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(["Ubicación", "Código artículo", "Descripción", "Cantidad", "Fecha/hora", "Usuario"])
    for f in filas:
        ws.append([
            f.get("ubicacion_codigo"), f.get("codigo_articulo"), f.get("descripcion"),
            f.get("cantidad"), f.get("creado_en"), f.get("usuario"),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=inventario_{desde}_a_{hasta}.xlsx"},
    )


@app.route("/api/inventario/<int:ubicacion_id>")
@auth.admin_required
def api_inventario_ubicacion(ubicacion_id):
    return jsonify(db.listar_inventario_ubicacion(ubicacion_id))


@app.route("/api/inventario/estado", methods=["GET", "POST"])
@auth.admin_required
def api_inventario_estado():
    # La pantalla de escaneo (PC) manda su estado actual acá cada vez que
    # cambia algo — sin esperar a que se guarde nada. La pantalla "en vivo"
    # (celular) lo consulta cada pocos segundos y lo muestra tal cual.
    if request.method == "POST":
        datos = request.get_json(silent=True) or {}
        db.actualizar_estado_vivo(
            usuario=session["user"],
            ubicacion_codigo=datos.get("ubicacion_codigo"),
            articulo_codigo=datos.get("articulo_codigo"),
            articulo_descripcion=datos.get("articulo_descripcion"),
            cantidad=datos.get("cantidad"),
            posicion=datos.get("posicion"),
            paso=datos.get("paso") or "",
        )
        return jsonify(ok=True)

    return jsonify(
        activos=db.listar_estado_vivo(),
        recientes=db.listar_inventario_reciente(30),
    )


@app.route("/admin/laboratorio/inventario/en-vivo")
@auth.admin_required
def admin_laboratorio_inventario_en_vivo():
    return render_template("admin_laboratorio_inventario_en_vivo.html")


@app.route("/admin/laboratorio/reconocimiento")
@auth.admin_required
def admin_laboratorio_reconocimiento():
    return render_template("admin_laboratorio_reconocimiento.html", cantidad_articulos=db.contar_articulos())


@app.route("/admin/laboratorio/reconocimiento/analizar", methods=["POST"])
@auth.admin_required
def admin_laboratorio_reconocimiento_analizar():
    # Banco de pruebas del reconocimiento: usa el mismo ocr_parser.parse_remito
    # que el flujo real de recepción, pero acá NO crea remito ni escribe nada
    # en la base. Si se eligió un proveedor, además busca cada ítem leído
    # contra el catálogo de artículos de ESE proveedor, para marcar si
    # existe y coincide con lo que dice el remito.
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        return jsonify(ok=False, error="Subí una foto o PDF."), 400
    if not _archivo_valido(archivo.filename):
        return jsonify(ok=False, error="Formato de archivo no admitido."), 400

    proveedor_id = request.form.get("proveedor_id", type=int)

    ext = Path(archivo.filename).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        archivo.save(tmp.name)
        tmp_path = tmp.name
    try:
        datos = ocr_parser.parse_remito(tmp_path, archivo.filename)
    finally:
        os.unlink(tmp_path)

    items = [
        {**it, **_matchear_articulo(it.get("codigo_articulo"), it.get("descripcion"), proveedor_id)}
        for it in datos["items"]
    ]

    return jsonify(
        ok=True,
        nro_remito=datos["nro_remito"],
        fecha_remito=datos["fecha_remito"],
        metodo_extraccion=datos["metodo_extraccion"],
        items=items,
        catalogo_consultado=proveedor_id is not None,
        cantidad_articulos_catalogo=db.contar_articulos(),
    )


@app.route("/admin/laboratorio/compras")
@auth.admin_required
def admin_laboratorio_compras():
    return render_template("admin_laboratorio_compras.html", compras=db.listar_compras_cargadas())


@app.route("/admin/laboratorio/compras/cargar", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_compras_cargar():
    # Misma lógica que "Nuevo remito" (subir foto/PDF y leer con Claude
    # Vision), pero para el documento de compra (factura/OC de proveedor),
    # y acepta también Excel — reutilizando el parser genérico de columnas
    # que ya usa el catálogo de artículos. Nunca se inventa una cantidad ni
    # un artículo: si no se reconoce nada, se avisa el error.
    if request.method == "GET":
        return render_template("admin_laboratorio_compras_cargar.html")

    proveedor_id = request.form.get("proveedor_id", type=int)
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        return render_template("admin_laboratorio_compras_cargar.html", error="Subí un archivo (foto, PDF o Excel)."), 400
    if not _archivo_compra_valido(archivo.filename):
        return render_template("admin_laboratorio_compras_cargar.html", error="Formato de archivo no admitido."), 400

    ext = Path(archivo.filename).suffix.lower()
    contenido = archivo.read()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(contenido)
        tmp_path = tmp.name

    nro_documento = None
    fecha_documento = None
    try:
        if ext in (".xlsx", ".xlsm", ".xls"):
            tipo_archivo = "excel"
            metodo = "excel"
            try:
                if _es_xlsx_real(tmp_path):
                    filas = _parsear_tabla_generica_xlsx(tmp_path, _CAMPOS_COMPRA_ITEMS, obligatorios=["codigo_articulo"])
                else:
                    filas = _parsear_tabla_generica_html(tmp_path, _CAMPOS_COMPRA_ITEMS, obligatorios=["codigo_articulo"])
            except Exception as e:
                os.unlink(tmp_path)
                return render_template("admin_laboratorio_compras_cargar.html", error=f"No se pudo leer el archivo: {e}"), 400

            items = []
            for fila in filas:
                codigo = str(fila["codigo_articulo"]).strip() if fila.get("codigo_articulo") else None
                if not codigo:
                    continue
                cantidad_num, cantidad_txt = _num_o_texto(fila.get("cantidad_comprada"))
                deposito_num, deposito_txt = _num_o_texto(fila.get("reparto_deposito_col"))
                # SOLANO = DELMY 1, VARELA = DELMY 3 (confirmado con Nicotoli).
                delmy1_num, delmy1_txt = _num_o_texto(fila.get("reparto_solano"))
                delmy3_num, delmy3_txt = _num_o_texto(fila.get("reparto_varela"))
                items.append({
                    "codigo_articulo": codigo,
                    "descripcion": str(fila["descripcion"]).strip() if fila.get("descripcion") else None,
                    "cantidad_comprada": cantidad_num,
                    "cantidad_comprada_texto": cantidad_txt,
                    "reparto_deposito": deposito_num,
                    "reparto_deposito_texto": deposito_txt,
                    "reparto_delmy1": delmy1_num,
                    "reparto_delmy1_texto": delmy1_txt,
                    "reparto_delmy3": delmy3_num,
                    "reparto_delmy3_texto": delmy3_txt,
                })
        else:
            tipo_archivo = "pdf" if ext == ".pdf" else "imagen"
            metodo = "claude_vision"
            datos = ocr_parser.parse_factura_compra(tmp_path, archivo.filename)
            items = datos["items"]
            nro_documento = datos["nro_documento"]
            fecha_documento = datos["fecha_documento"]
    finally:
        os.unlink(tmp_path)

    if not items:
        return render_template(
            "admin_laboratorio_compras_cargar.html",
            error="No se reconoció ningún artículo en el archivo. Probá con otro escaneo, o cargá el detalle a mano más adelante.",
        ), 400

    compra_id = db.crear_compra_cargada(
        proveedor_id=proveedor_id,
        archivo_original=None,
        tipo_archivo=tipo_archivo,
        metodo_extraccion=metodo,
        nro_documento=nro_documento,
        fecha_documento=fecha_documento,
        usuario_creador=session["user"],
        items=items,
    )

    nombre_final = f"compra_{compra_id}_{secure_filename(Path(archivo.filename).stem)}{ext}"
    tipo_mime = archivo.mimetype or mimetypes.guess_type(nombre_final)[0] or "application/octet-stream"
    db.guardar_archivo(nombre_final, contenido, tipo_mime)
    db.actualizar_archivo_compra(compra_id, nombre_final)

    return redirect(url_for("admin_laboratorio_compras_detalle", compra_id=compra_id))


@app.route("/admin/laboratorio/compras/<int:compra_id>", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_compras_detalle(compra_id):
    compra = db.get_compra_cargada(compra_id)
    if not compra:
        return "Compra no encontrada", 404

    if request.method == "POST":
        # Guarda el reparto entre depósito y sucursales — esto SIEMPRE lo
        # completa una persona a mano: no se lee de ningún documento porque
        # normalmente no está escrito ahí.
        items = db.get_compra_items(compra_id)
        repartos = {}
        for it in items:
            def _num(prefijo):
                val = request.form.get(f"{prefijo}_{it['id']}", "").strip()
                try:
                    return float(val) if val else None
                except ValueError:
                    return None
            repartos[it["id"]] = {
                "deposito": _num("reparto_deposito"),
                "delmy1": _num("reparto_delmy1"),
                "delmy3": _num("reparto_delmy3"),
            }
        db.actualizar_reparto_compra_items(compra_id, repartos)
        return redirect(url_for("admin_laboratorio_compras_detalle", compra_id=compra_id))

    return render_template(
        "admin_laboratorio_compras_detalle.html",
        compra=compra,
        items=db.get_compra_items(compra_id),
    )


# ------------------------------------------------------- ubicaciones (admin)
#
# Etiquetas de posición física del depósito (pasillo/estantería/altura/
# ubicación), para pegar en las estanterías y escanear con lector de código
# de barras o el celular. Reutiliza el mismo generador de PDF (PyMuPDF) que
# las etiquetas de remito, agregando el código de barras Code128.

@app.route("/admin/laboratorio/ubicaciones")
@auth.admin_required
def admin_laboratorio_ubicaciones():
    return render_template(
        "admin_laboratorio_ubicaciones.html",
        ubicaciones=db.listar_ubicaciones(),
        cola_ids={u["id"] for u in db.listar_cola_impresion()},
    )


@app.route("/admin/laboratorio/ubicaciones/agregar", methods=["POST"])
@auth.admin_required
def admin_laboratorio_ubicaciones_agregar():
    # Alta rápida sin recargar la página: el código de cada nivel es el que
    # se escribe tal cual (nada de ×10 automático) — así se puede cargar
    # tanto el patrón habitual (10/20/30, 15/25/35...) como separaciones
    # especiales (12, 14, 17, 19...) sin que la app imponga un número.
    datos = request.get_json(silent=True) or {}
    pasillo = (datos.get("pasillo") or "").strip()
    estanteria = (datos.get("estanteria") or "").strip() or None
    altura = (datos.get("altura") or "").strip() or None
    ubicacion = (datos.get("ubicacion") or "").strip() or None
    descripcion = (datos.get("descripcion") or "").strip() or None

    if not pasillo:
        return jsonify(ok=False, error="Falta el pasillo."), 400
    if altura and not estanteria:
        return jsonify(ok=False, error="Para cargar la altura primero hace falta la estantería."), 400
    if ubicacion and not altura:
        return jsonify(ok=False, error="Para cargar la ubicación primero hace falta la altura."), 400

    ubicacion_creada, error = db.crear_ubicacion(pasillo, estanteria, altura, ubicacion, None, descripcion)
    if error:
        return jsonify(ok=False, error=error), 400
    return jsonify(ok=True, ubicacion=ubicacion_creada)


@app.route("/admin/laboratorio/ubicaciones/<int:ubicacion_id>/eliminar", methods=["POST"])
@auth.admin_required
def admin_laboratorio_ubicaciones_eliminar(ubicacion_id):
    db.eliminar_ubicacion(ubicacion_id)
    return jsonify(ok=True)


@app.route("/admin/laboratorio/ubicaciones/eliminar-multiples", methods=["POST"])
@auth.admin_required
def admin_laboratorio_ubicaciones_eliminar_multiples():
    datos = request.get_json(silent=True) or {}
    ids = [int(i) for i in (datos.get("ids") or [])]
    borradas, bloqueadas = db.eliminar_ubicaciones(ids)
    return jsonify(ok=True, borradas=borradas, bloqueadas=bloqueadas)


@app.route("/admin/laboratorio/ubicaciones/eliminar-todas", methods=["POST"])
@auth.admin_required
def admin_laboratorio_ubicaciones_eliminar_todas():
    borradas, bloqueadas = db.eliminar_todas_ubicaciones()
    return jsonify(ok=True, borradas=borradas, bloqueadas=bloqueadas)


@app.route("/admin/laboratorio/ubicaciones/cola", methods=["GET", "POST"])
@auth.admin_required
def admin_laboratorio_ubicaciones_cola():
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "agregar":
            ids = [int(i) for i in request.form.getlist("ubicacion_id")]
            db.agregar_a_cola_impresion(ids)
        elif accion == "quitar":
            db.quitar_de_cola_impresion(request.form.get("ubicacion_id", type=int))
        elif accion == "vaciar":
            db.vaciar_cola_impresion()
        return redirect(request.referrer or url_for("admin_laboratorio_ubicaciones_cola"))

    return render_template("admin_laboratorio_ubicaciones_cola.html", cola=db.listar_cola_impresion())


@app.route("/admin/laboratorio/ubicaciones/etiquetas")
@auth.admin_required
def admin_laboratorio_ubicaciones_etiquetas():
    # Vista previa antes de imprimir — mismo patrón que las etiquetas de
    # remito: grilla de etiquetas en pantalla, con botones para disparar el
    # diálogo de impresión (WiFi/red) o abrir el PDF directo (USB/Xprinter).
    ids_param = request.args.get("ids")
    if ids_param:
        ids = [int(i) for i in ids_param.split(",") if i.strip().isdigit()]
        ubicaciones = db.get_ubicaciones_por_ids(ids)
    else:
        ubicaciones = db.listar_cola_impresion()
    return render_template(
        "admin_laboratorio_ubicaciones_etiquetas.html",
        ubicaciones=ubicaciones,
        ids=ids_param or "",
    )


@app.route("/admin/laboratorio/ubicaciones/etiquetas/imprimir")
@auth.admin_required
def admin_laboratorio_ubicaciones_imprimir():
    pdf_url = url_for("admin_laboratorio_ubicaciones_pdf", **request.args)
    return render_template("imprimir_etiquetas.html", pdf_url=pdf_url)


@app.route("/admin/laboratorio/ubicaciones/etiquetas.pdf")
@auth.admin_required
def admin_laboratorio_ubicaciones_pdf():
    # ?ids=1,2,3 imprime solo esas (ej. "imprimir esta sola" desde el
    # listado); sin parámetro, imprime toda la cola actual.
    ids_param = request.args.get("ids")
    if ids_param:
        ids = [int(i) for i in ids_param.split(",") if i.strip().isdigit()]
        ubicaciones = db.get_ubicaciones_por_ids(ids)
    else:
        ubicaciones = db.listar_cola_impresion()

    codigos = [u["codigo"] for u in ubicaciones]
    pdf_bytes = generar_pdf_etiquetas_ubicaciones(
        codigos,
        ancho_mm=request.args.get("ancho", 55, type=float),
        alto_mm=request.args.get("alto", 45, type=float),
    )
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": "inline; filename=etiquetas_ubicaciones.pdf"},
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=False, host="0.0.0.0", port=port)
